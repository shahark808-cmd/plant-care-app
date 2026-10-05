"""מוטי — סוכן איתור לקוחות לתדיראן טלקום, חינמי. מחקר ודירוג בלבד, בלי פנייה ללקוחות.

שימוש: python agent.py [--sector "ביטוח ופנסיה"] [--limit 5] [--no-email]
עם SUPABASE_URL/SUPABASE_SERVICE_KEY כותב ל-Supabase; בלעדיהם כותב קובץ Excel מקומי.
"""
import argparse, datetime, json, pathlib, sys
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

import notify, persona, score, store, verify
from sources import jobs, maya, tenders

HERE = pathlib.Path(__file__).parent
SEEN = HERE / "seen.json"
COLS = [("name", "חברה"), ("sector", "ענף"), ("employees_estimate", "עובדים (משוער)"),
        ("employees_confidence", "ביטחון בגודל"), ("score", "ציון"), ("rationale", "נימוק"),
        ("pain_signals", "סימני כאב"), ("channels_today", "ערוצים כיום"),
        ("contact_center_notes", "מוקד"), ("contact_role", "איש קשר מוצע"), ("sources", "מקורות")]


def in_range(lead, icp):
    n = lead.get("employees_estimate")
    return n is None or icp["employees_min"] <= n <= icp["employees_max"]


def export_excel(leads):
    wb = Workbook(); ws = wb.active; ws.title = "לידים"; ws.sheet_view.rightToLeft = True
    ws.append([h for _, h in COLS])
    for c in ws[1]:
        c.font = Font(bold=True)
    for l in sorted(leads, key=lambda x: -(x.get("score") or 0)):
        ws.append([", ".join(v) if isinstance(v := l.get(k), list) else v for k, _ in COLS])
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    path = HERE / f"leads_{datetime.date.today()}.xlsx"
    wb.save(path)
    return path


def run(icp, sectors, send_email=True):
    if store.enabled():
        seen, reasons = store.known()
    else:
        seen = set(json.loads(SEEN.read_text(encoding="utf-8"))) if SEEN.exists() else set()
        reasons = []
    leads = []
    for sector in sectors:
        print(f"{persona.NAME}: יוצא לחפש בענף {sector}...", file=sys.stderr)
        ctx = "\n".join((jobs.discover_text(sector), tenders.discover_text(sector), maya.discover_text(sector)))
        for cand in score.discover_candidates(icp, sector, ctx, seen):
            name = (cand.get("name") or "").strip()
            if not name or name in seen:
                continue
            seen.add(name)
            ev = verify.gather(cand)
            if ev["registry"] and not ev["registry"]["active"]:
                print(f"  - {name}: לא פעילה", file=sys.stderr); continue
            lead = score.rate(icp, cand, sector, ev["text"], reasons)
            if not lead or not in_range(lead, icp) or (lead.get("score") or 0) < icp["min_score"]:
                print(f"  - {name}: נדחתה", file=sys.stderr); continue
            lead["name"] = name
            print(f"  + {name}: {lead['score']}", file=sys.stderr)
            leads.append(lead)
    if store.enabled():
        print(f"נשמרו ל-Supabase: {store.save(leads)}")
    else:
        SEEN.write_text(json.dumps(sorted(seen), ensure_ascii=False), encoding="utf-8")
        print(f"נשמר: {export_excel(leads)} ({len(leads)} לידים)")
    if send_email and leads:
        print("מייל נשלח" if notify.send(leads, icp.get("hot_score", 80)) else "מייל לא הוגדר, דילוג")
    return leads


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sector"); ap.add_argument("--limit", type=int); ap.add_argument("--no-email", action="store_true")
    a = ap.parse_args()
    icp = yaml.safe_load((HERE / "icp.yaml").read_text(encoding="utf-8"))
    if a.limit:
        icp["per_sector_candidates"] = a.limit
    run(icp, [a.sector] if a.sector else icp["sectors"], not a.no_email)


if __name__ == "__main__":
    main()
